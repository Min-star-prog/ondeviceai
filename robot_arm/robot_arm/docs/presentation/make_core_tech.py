"""Build an editable Core Tech section from the current robot-arm software."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ASSET = HERE / "assets"
ROOT = HERE.parents[1]
OUT = HERE / "robot_arm_core_tech.pptx"

prs = Presentation()
prs.slide_width = Inches(20)
prs.slide_height = Inches(11.25)
BLANK = prs.slide_layouts[6]

BLUE = RGBColor(39, 84, 188)
NAVY = RGBColor(15, 49, 111)
INK = RGBColor(19, 26, 36)
MUTED = RGBColor(91, 108, 127)
LINE = RGBColor(211, 220, 232)
PALE = RGBColor(238, 244, 255)
WHITE = RGBColor(255, 255, 255)
ORANGE = RGBColor(239, 106, 65)
GREEN = RGBColor(66, 151, 109)


def shape(slide, x, y, w, h, fill=WHITE, stroke=LINE, radius=False, width=1.1):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = stroke
    s.line.width = Pt(width)
    if radius:
        try: s.adjustments[0] = 0.12
        except Exception: pass
    return s


def text(slide, value, x, y, w, h, size=23, bold=False, color=INK,
         align=PP_ALIGN.LEFT, font="Malgun Gothic", margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        run = p.runs[0]
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        rpr = run._r.get_or_add_rPr()
        ea = OxmlElement("a:ea")
        ea.set("typeface", font)
        rpr.append(ea)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
    return box


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def image(slide, path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def new_slide(title, num):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.82), Inches(0.91), Inches(0.48), Inches(0.48))
    circle.fill.solid(); circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    text(s, "›", 0.82, 0.88, 0.48, 0.48, 30, True, WHITE, PP_ALIGN.CENTER)
    text(s, title, 1.48, 0.61, 14.8, 1.06, 49, True)
    image(s, ASSET / "kcci_logo.png", 16.27, 0.34, 3.1, 0.56)
    line(s, 0.82, 2.0, 19.17, 2.0, LINE, 1.3)
    text(s, "2. Core Tech", 0.84, 10.77, 3.0, 0.27, 13, False, MUTED)
    text(s, f"{num:02d}", 18.55, 10.76, 0.56, 0.3, 13, False, MUTED, PP_ALIGN.RIGHT)
    return s


def card(slide, x, y, w, h, label, body, accent=BLUE, body_size=23):
    shape(slide, x, y, w, h, WHITE, LINE, True)
    shape(slide, x, y, 0.12, h, accent, accent)
    text(slide, label, x+0.31, y+0.25, w-0.56, 0.53, 25, True, accent)
    text(slide, body, x+0.31, y+0.85, w-0.56, h-1.08, body_size, False, INK)


# 1 — Section cover.
s = prs.slides.add_slide(BLANK)
s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
shape(s, 0, 0, 20, 0.22, NAVY, NAVY)
shape(s, 1.05, 3.02, 0.17, 4.9, BLUE, BLUE)
text(s, "02", 1.67, 2.86, 2.4, 1.2, 55, True, BLUE)
text(s, "Core Tech", 1.64, 4.02, 12.1, 1.25, 69, True, INK)
text(s, "사람 동작 해석 · 5축 로봇 제어 · 기록/재생", 1.69, 5.54, 15.3, 0.77, 29, False, MUTED)
image(s, ASSET / "kcci_logo.png", 16.1, 0.52, 3.15, 0.57)
text(s, "로봇팔 제어 소프트웨어", 1.68, 8.67, 9.0, 0.52, 23, True, NAVY)

# 2 — Architecture: reuse the editable SW stack diagram's exported body.
s = new_slide("로봇팔 제어 SW 구조", 2)
image(s, ASSET / "sw_stack_body.png", 0.92, 2.27, 18.12, 7.71)

# 3 — Human posture interpretation.
s = new_slide("사람 자세 해석", 3)
shape(s, 0.82, 2.57, 9.0, 5.36, WHITE, LINE)
image(s, ROOT / "docs/diagrams/robot_overview_frame.jpg", 0.91, 2.66, 8.82, 5.18)
text(s, "실제 영상 입력과 3D 자세를 함께 확인한 PC 재생 화면", 0.88, 8.09, 9.1, 0.45, 19, False, MUTED)
steps = [
    ("01", "2D 자세 입력", "양어깨·팔꿈치·손목·손가락 좌표"),
    ("02", "3D 복원과 EMA", "상대 깊이를 복원하고 좌표 흔들림 완화"),
    ("03", "몸 기준 좌표계", "양어깨와 CameraUp으로 BodyFrame 계산"),
    ("04", "사람 목표값", "팔꿈치·손목 4각도와 그리퍼 개폐값"),
]
for i,(n,label,body) in enumerate(steps):
    y=2.51+i*1.64
    shape(s, 10.55, y, 8.62, 1.38, PALE, LINE, True)
    text(s, n, 10.8, y+0.28, 0.75, 0.68, 29, True, BLUE, PP_ALIGN.CENTER)
    text(s, label, 11.76, y+0.14, 7.04, 0.51, 25, True)
    text(s, body, 11.76, y+0.73, 7.04, 0.43, 20, False, MUTED)

# 4 — Mapping and trajectory.
s = new_slide("사람 각도에서 로봇 동작까지", 4)
shape(s, 1.0, 2.44, 18.0, 1.18, PALE, RGBColor(185,204,236), True)
text(s, "서보 목표각 = 사람 각도 × 크기 × 방향 + 기준각", 1.52, 2.61, 16.96, 0.81,
     35, True, NAVY, PP_ALIGN.CENTER)
card(s, 1.0, 4.0, 5.65, 3.65, "각도 보정", "관절별 방향과 기준각을 적용\n예: 팔꿈치 방위각 30° → 서보 60°", BLUE, 23)
card(s, 7.18, 4.0, 5.65, 3.65, "명령 검사", "서보 범위 확인\n3D FK로 자기 충돌·테이블 충돌 검사", ORANGE, 23)
card(s, 13.36, 4.0, 5.65, 3.65, "움직임 계산", "4개 회전축에 속도·가속도 제한 적용\n다음 명령 각도를 함께 갱신", GREEN, 23)
shape(s, 1.0, 8.13, 18.0, 1.37, RGBColor(247,249,253), LINE, True)
text(s, "입력 프레임이 잠시 끊겨도 마지막으로 승인한 목표까지 이동한 뒤 명령 자세를 유지", 1.45, 8.42,
     17.1, 0.74, 25, True, NAVY, PP_ALIGN.CENTER)

# 5 — PWM channels and the actual register contract.
s = new_slide("5축 서보 PWM 출력", 5)
text(s, "출력 채널", 1.0, 2.51, 8.2, 0.53, 26, True, NAVY)
channel_rows = [
    ("CH0", "팔꿈치 방위각", "elbow_roll"),
    ("CH1", "팔꿈치 고도각", "elbow_pitch"),
    ("CH2", "손목 굽힘", "wrist_pitch"),
    ("CH3", "손목 회전", "wrist_roll"),
    ("CH4", "그리퍼", "gripper"),
]
for i,(ch,kr,en) in enumerate(channel_rows):
    y=3.2+i*1.13
    shape(s, 1.0, y, 8.1, 0.94, PALE if i%2==0 else WHITE, LINE)
    text(s, ch, 1.28, y+0.17, 1.17, 0.54, 24, True, BLUE)
    text(s, kr, 2.53, y+0.17, 3.54, 0.54, 23, True)
    text(s, en, 6.08, y+0.18, 2.75, 0.52, 18, False, MUTED)
shape(s, 10.1, 2.54, 8.9, 6.42, WHITE, LINE, True)
text(s, "출력 단계", 10.55, 2.8, 7.93, 0.64, 28, True, NAVY)
for i,(label,detail) in enumerate([
    ("각도 → 펄스폭", "축별 min·center·max 기준값으로 µs 변환"),
    ("레지스터 기록", "CH0~CH4 값을 AXI4-Lite에 기록"),
    ("동시 적용", "UPDATE 레지스터로 5채널 값 반영"),
]):
    y=3.69+i*1.56
    shape(s, 10.55, y, 7.99, 1.26, RGBColor(247,249,253), LINE, True)
    text(s, f"0{i+1}", 10.81, y+0.2, 0.67, 0.76, 29, True, BLUE)
    text(s, label, 11.62, y+0.1, 6.51, 0.48, 23, True)
    text(s, detail, 11.62, y+0.59, 6.52, 0.44, 18, False, MUTED)
text(s, "기본 통합 빌드에서는 로봇 PWM 적용을 끄고, 별도 빌드에서 출력합니다.",
     1.0, 9.45, 17.9, 0.54, 18, False, MUTED)

# 6 — Record/replay controller.
s = new_slide("동작 기록·재생", 6)
shape(s, 1.0, 2.55, 18.0, 1.0, RGBColor(248,250,253), LINE, True)
text(s, "녹화", 1.31, 2.76, 1.31, 0.57, 25, True, BLUE)
text(s, "LIVE  ⇄  RECORDING", 3.36, 2.73, 9.2, 0.62, 28, True, NAVY)
text(s, "UART R로 시작·종료", 13.03, 2.77, 5.46, 0.52, 20, False, MUTED)
shape(s, 1.0, 3.76, 18.0, 1.0, RGBColor(245,250,247), LINE, True)
text(s, "재생", 1.31, 3.97, 1.31, 0.57, 25, True, GREEN)
text(s, "LIVE  →  ALIGNING  →  PLAYING  →  HOLDING", 3.36, 3.95, 12.3, 0.62, 27, True, GREEN)
text(s, "P로 LIVE 복귀", 15.72, 3.99, 2.85, 0.52, 19, False, MUTED)
card(s, 0.98, 5.2, 5.65, 3.63, "기록", "검증·출력한 5축 명령을 저장\n20 ms 간격, 최대 1024샘플", BLUE, 23)
card(s, 7.17, 5.2, 5.65, 3.63, "재생 전 검사", "저장된 전체 명령의 범위·FK·속도·가속도 확인", ORANGE, 23)
card(s, 13.36, 5.2, 5.65, 3.63, "1회 재생", "첫 명령까지 정렬한 뒤 순서대로 출력\n완료 후 마지막 명령 유지", GREEN, 23)

# 7 — Evidence from current CSV and matching PC viewer.
s = new_slide("실제 영상으로 자세 해석 검증", 7)
metric = [("853", "입력 프레임"), ("737", "유효한 사람 목표값"), ("42.6초", "영상 구간")]
for i,(n,label) in enumerate(metric):
    x=1.04+i*6.07
    shape(s, x, 2.34, 5.62, 1.14, PALE, LINE, True)
    text(s, n, x+0.23, 2.55, 1.82, 0.66, 33, True, NAVY)
    text(s, label, x+2.06, 2.57, 3.31, 0.62, 21, False, MUTED)
shape(s, 0.99, 3.9, 8.55, 5.0, WHITE, LINE)
image(s, ROOT / "docs/diagrams/robot_overview_frame.jpg", 1.08, 3.99, 8.37, 4.82)
image(s, ASSET / "example6_angles.png", 9.9, 4.15, 9.1, 4.31)
text(s, "PC 영상·CSV 재생에서 계산한 사람 목표각입니다. 실제 서보 위치 측정값은 아닙니다.",
     1.0, 9.22, 18.0, 0.58, 19, False, MUTED)

prs.save(OUT)
print(OUT, len(prs.slides))
